#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# 灵思微致测试平台 · 文档渲染器(V2.5 轻量版,确定性,不调 LLM)
# 用法: python3 render_doc.py <pptx|docx|pdf> <spec.json> <out_path>
# 退出码 0=成功,stdout 输出产物摘要(页数/段落数);非 0=失败,stderr 原因
import json
import sys


def load_spec(path):
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def render_pptx(spec, out):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    slides = spec.get('slides') or []
    if not slides:
        raise ValueError('规格缺少 slides')
    blank = prs.slide_layouts[6]
    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        title_text = str(s.get('title') or '')
        bullets = s.get('bullets') or []
        subtitle = str(s.get('subtitle') or '')
        # 标题文本框
        tb = slide.shapes.add_textbox(Pt(30), Pt(30), Pt(680), Pt(80))
        tf = tb.text_frame
        tf.text = title_text
        tf.paragraphs[0].runs[0].font.size = Pt(32 if i == 0 else 24)
        tf.paragraphs[0].runs[0].font.bold = True
        y = Pt(120)
        if subtitle:
            sb = slide.shapes.add_textbox(Pt(30), y, Pt(680), Pt(50))
            sb.text_frame.text = subtitle
            sb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
            y = Pt(180)
        if bullets:
            bb = slide.shapes.add_textbox(Pt(40), y, Pt(660), Pt(360))
            btf = bb.text_frame
            btf.word_wrap = True
            for j, b in enumerate(bullets):
                p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
                p.text = '• ' + str(b)
                p.runs[0].font.size = Pt(16)
                p.space_after = Pt(8)
    prs.save(out)
    return '%d 页幻灯片' % len(slides)


def render_docx(spec, out):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(11)
    title = str(spec.get('title') or '文档')
    doc.add_heading(title, level=0)
    blocks = spec.get('blocks') or []
    if not blocks:
        raise ValueError('规格缺少 blocks')
    n = 0
    for b in blocks:
        t = str(b.get('type') or 'p')
        text = str(b.get('text') or '')
        if not text:
            continue
        if t == 'h1':
            doc.add_heading(text, level=1)
        elif t == 'h2':
            doc.add_heading(text, level=2)
        elif t == 'bullet':
            doc.add_paragraph(text, style='List Bullet')
        else:
            doc.add_paragraph(text)
        n += 1
    doc.save(out)
    return '%d 个内容块' % n


def render_pdf(spec, out):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))  # 内置 CID 中文,无需字体文件
    c = canvas.Canvas(out, pagesize=A4)
    width, height = A4
    x, y = 20 * mm, height - 25 * mm
    title = str(spec.get('title') or '文档')
    c.setFont('STSong-Light', 20)
    c.drawString(x, y, title)
    y -= 15 * mm
    blocks = spec.get('blocks') or []
    if not blocks:
        raise ValueError('规格缺少 blocks')
    pages = 1
    for b in blocks:
        t = str(b.get('type') or 'p')
        text = str(b.get('text') or '')
        if not text:
            continue
        size = 15 if t == 'h1' else 13 if t == 'h2' else 10.5
        c.setFont('STSong-Light', size)
        prefix = '• ' if t == 'bullet' else ''
        # 简易换行:按页宽 45 个中文字符折行
        line = prefix + text
        while line:
            seg, line = line[:45], line[45:]
            if y < 20 * mm:
                c.showPage()
                c.setFont('STSong-Light', size)
                y = height - 25 * mm
                pages += 1
            c.drawString(x, y, seg)
            y -= 8 * mm
        y -= 2 * mm
    c.save()
    return '%d 页 PDF' % pages


def main():
    if len(sys.argv) != 4:
        print('用法: render_doc.py <pptx|docx|pdf> <spec.json> <out>', file=sys.stderr)
        return 2
    kind, spec_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
    spec = load_spec(spec_path)
    if kind == 'pptx':
        msg = render_pptx(spec, out)
    elif kind == 'docx':
        msg = render_docx(spec, out)
    elif kind == 'pdf':
        msg = render_pdf(spec, out)
    else:
        print('未知类型: ' + kind, file=sys.stderr)
        return 2
    print(msg)
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as e:
        print('渲染异常: %s' % e, file=sys.stderr)
        sys.exit(1)
